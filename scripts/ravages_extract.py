"""Extract the army movements (ComLoc) and what explains them from the Ravages
master sheet into one JSON for the army tracker page.

    python -I scripts/ravages_extract.py <Ravages_Master_sheet.xlsx> ravages/data.json
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
NCOL = 90
EVENT_KINDS = [  # (kind, pattern) checked in order; anything else is routine and dropped
    ("battle", r"battle|victory|defeat|routed|ambush"),
    ("assault", r"assault|siege|storm|sack"),
    ("operation", r"operation"),
    ("march", r"force[d]? march|rest\b"),
    ("command", r"subcommander|command|joined|merge|split"),
    ("recruit", r"recruit"),
    ("tax", r"taxation|tax collection"),
]


def iso(d):
    return d.date().isoformat() if isinstance(d, dt.datetime) else None


def main(xlsx, out):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    loc = list(wb["ComLoc"].iter_rows(values_only=True))
    names = [str(n).strip() for n in loc[0][2:] if n]
    dates, pos = [], {n: [] for n in names}
    for r in loc[1:]:
        if not isinstance(r[0], dt.datetime):
            break
        dates.append(iso(r[0]))
        for n, v in zip(names, r[2:2 + len(names)]):
            pos[n].append(int(v) if isinstance(v, (int, float)) else None)
    dtab = {iso(r[0]): (r[1], r[2]) for r in list(wb["DateTable"].iter_rows(values_only=True))[1:] if r and r[0]}
    days = [dict(date=d, ig=(dtab.get(d) or (None,))[0], season=(dtab.get(d) or (None, None))[1]) for d in dates]

    com = list(wb["Commanders"].iter_rows(values_only=True))
    cnames = [str(c).strip() if c else None for c in com[0]]
    label_rows = {str(r[0]).strip(): r for r in com if r and r[0]}
    numbers_at = next(i for i, r in enumerate(com) if r and r[0] == "Numbers")
    units = [str(r[0]).strip() for r in com[numbers_at + 1:numbers_at + 8]]
    info = {}
    for j, n in enumerate(cnames):
        if not n or j == 0:
            continue
        g = lambda lab: (label_rows.get(lab) or [None] * (j + 1))[j]
        troops = {u: round(float(com[numbers_at + 1 + k][j] or 0)) for k, u in enumerate(units)}
        info[n] = dict(traits=[t for t in (g("Trait 1"), g("Trait 2"), g("Trait 3")) if t],
                       tradition=g("Army Tradition"), troops={k: v for k, v in troops.items() if v},
                       morale=g("Current Morale"), speed=g("Speed (Miles/Day)"),
                       supplies=round(float(g("Current Supplies") or 0)), coin=round(float(g("Current Coin") or 0)))

    events = []
    for r in list(wb["ComEvents"].iter_rows(values_only=True))[1:]:
        if not r or not r[0] or not isinstance(r[1], dt.datetime) or not r[5]:
            continue
        note = str(r[5]).strip()
        kind = next((k for k, pat in EVENT_KINDS if re.search(pat, note, re.I)), None)
        if kind is None:
            continue
        events.append(dict(army=str(r[0]).strip(), date=iso(r[1]), kind=kind, note=note,
                           morale=r[4] if isinstance(r[4], (int, float)) else None))

    hexes = {}
    for r in list(wb["Hex"].iter_rows(values_only=True))[2:]:
        if r[0] is None:
            continue
        h = int(r[0])
        row, col = int(r[20]), int(r[21])
        assert h == (row - 1) * NCOL + col, (h, row, col)
        terrain = r[5]
        if terrain in (None, "N/A", "Ocean", "Sea"):
            continue
        hexes[h] = [r[2] if r[2] not in (None, "N/A") else "", r[3] if r[3] not in (None, "N/A") else "",
                    r[4] or "", terrain, r[11] or 0]
    towns = []
    for r in list(wb["TownsCities"].iter_rows(values_only=True))[1:]:
        if r and isinstance(r[0], (int, float)) and r[1]:
            towns.append([int(r[0]), str(r[1]).strip(), r[2] or "", r[6] or ""])
    grid = json.loads((ROOT / "data" / "hexgrid.json").read_text())
    data = dict(days=days, armies=[dict(name=n, hex=pos[n], info=info.get(n)) for n in names],
                events=events, hexes=hexes, towns=towns,
                grid=dict(ncol=NCOL, x0=grid["x0"], dx=grid["dx"], y0=grid["y0"], dy=grid["dy"], width=2000, height=926),
                source=Path(xlsx).name)
    Path(out).write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    kinds = {}
    for e in events:
        kinds[e["kind"]] = kinds.get(e["kind"], 0) + 1
    print(f"{len(names)} armies x {len(dates)} days ({dates[0]} .. {dates[-1]}), {len(events)} events {kinds}, "
          f"{len(hexes)} land hexes, {len(towns)} towns; {Path(out).stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
