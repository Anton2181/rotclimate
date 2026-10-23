"""Roman towns of AD 150 against the settlement model (and HYDE).

    python3 -I scripts/roman_cities_check.py --xlsx data/cache/hanson2016/cities.xlsx \
        --earth output/land_people/earth/earth.npz --hyde data/cache/hyde33/popc_ce_200.tif \
        --out output/population_roman/ad150.json

Towns: Hanson (2016), An Urban Geography of the Roman World, 100 BC to AD
300, OxREP cities database (https://oxrep.web.ox.ac.uk/cities-database):
1,388 sites with urban characteristics, their coordinates, dates and, for
885, an inhabited or walled area in hectares. The sites standing in AD 150
(start <= 150, end >= 150 or open) are counted per 2-degree block of land
over the empire's footprint (blocks with a site within 2 degrees). Two
targets per block: the number of towns, and their summed area (towns with
no area take the median area of their Barrington Atlas rank).

Town density follows rural density (towns lived on their hinterland's
surplus) but also the age of urban life (Greece and Asia Minor against
Gaul and Britain), so this tests the same thing the Han census tests, inside
another empire and with an independent source. Each model field is summed
per block; rank correlation of per-km2 values over the footprint, and the
share of towns (or area) that would have to move.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "scripts"))
from harvest_earth_hyde import earth_stages  # noqa: E402
from hyde_compare import blocks, hyde_at, spearman  # noqa: E402
from population_cores import CORE_SETS, core_distance, core_term, dry_river_cores, travel_distance  # noqa: E402

YEAR = 150
BLOCK = 24          # 5' cells: 2 degrees


def towns(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Cities"].iter_rows(values_only=True))
    head = rows[0]
    cities = [dict(zip(head, r)) for r in rows[1:] if r[0]]
    area = {r[0]: r[1] for r in list(wb["Areas"].iter_rows(values_only=True))[1:] if r[0]}

    def num(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    out = []
    for c in cities:
        s, e = num(c["Start Date"]), num(c["End Date"])
        lon, lat = num(c["Longitude (X)"]), num(c["Latitude (Y)"])
        if lon is None or lat is None or s is None or s > YEAR or (e is not None and e < YEAR):
            continue
        out.append(dict(key=c["Primary Key"], name=c["Ancient Toponym"], lon=lon, lat=lat,
                        rank=str(c["Barrington Atlas Rank"]), area_ha=num(area.get(c["Primary Key"]))))
    by_rank = {}
    for t in out:
        if t["area_ha"]:
            by_rank.setdefault(t["rank"], []).append(t["area_ha"])
    med_all = float(np.median([t["area_ha"] for t in out if t["area_ha"]]))
    for t in out:
        if not t["area_ha"]:
            t["area_ha"] = float(np.median(by_rank[t["rank"]])) if t["rank"] in by_rank else med_all
            t["area_filled"] = True
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--earth", required=True)
    ap.add_argument("--hyde")
    ap.add_argument("--out", required=True)
    ap.add_argument("--reliability", help="START,FULL,POWER of harvest.models.harvest_reliability (default none)")
    a = ap.parse_args()
    z = np.load(a.earth)
    lat, lon, land = z["lat"], z["lon"], z["land"]
    km2, stages = earth_stages(z, reliability=tuple(map(float, a.reliability.split(","))) if a.reliability else None)
    base = np.nan_to_num(stages["+ malaria (settlement base)"])
    cores = CORE_SETS["cradles + Rome"]
    fields = {
        "uniform over land": np.ones_like(base),
        "settlement base": base,
        "base x cores (cradles + Rome, great circle, L 800)": base * core_term(core_distance(lat, lon, cores), 800),
        "base x cores (travel: forest x2, relief 1/km, L 800)": base * core_term(travel_distance(z, cores, 2.0, 1.0, 1.0), 800),
    }
    rule = dry_river_cores(z, base, km2, 10)
    fields["base x rule cores (top 10, great circle, L 800)"] = base * core_term(core_distance(lat, lon, rule), 800)
    fields["base x rule cores (top 10, travel, L 800)"] = base * core_term(travel_distance(z, rule, 2.0, 1.0, 1.0), 800)
    if a.hyde:
        fields["HYDE 3.3, 200 CE"] = hyde_at(a.hyde, lat, lon) / km2
    T = towns(a.xlsx)
    ny, nx = land.shape[0] // BLOCK, land.shape[1] // BLOCK
    blat = lat[:ny * BLOCK].reshape(ny, BLOCK).mean(1)
    blon = lon[:nx * BLOCK].reshape(nx, BLOCK).mean(1)
    count = np.zeros((ny, nx)); area = np.zeros((ny, nx))
    step_lat, step_lon = lat[0] - lat[1], lon[1] - lon[0]
    used = 0
    for t in T:
        i = int((lat[0] + step_lat / 2 - t["lat"]) / step_lat) // BLOCK
        j = int((t["lon"] - lon[0] + step_lon / 2) / step_lon) // BLOCK
        if 0 <= i < ny and 0 <= j < nx:
            count[i, j] += 1; area[i, j] += t["area_ha"]; used += 1
    A = blocks(km2 * land, BLOCK)
    landshare = blocks(land.astype(float), BLOCK) / BLOCK ** 2
    near = np.zeros((ny, nx), bool)
    for i, j in zip(*np.nonzero(count)):
        near[max(0, i - 1):i + 2, max(0, j - 1):j + 2] = True
    ok = near & (landshare >= 0.3)
    res = dict(year=YEAR, towns_standing=len(T), towns_in_window=used, blocks=int(ok.sum()),
               towns_in_blocks=int(count[ok].sum()), stages={})
    print(f"AD {YEAR}: {len(T)} towns standing, {used} in the Earth window; {int(ok.sum())} land blocks of 2 degrees "
          f"over the footprint ({int(count[ok].sum())} towns)")
    for k, f in fields.items():
        M = blocks(np.nan_to_num(f) * km2 * land, BLOCK)[ok]
        v = {}
        for tname, tgt in (("towns", count[ok]), ("town area", area[ok])):
            m = M * tgt.sum() / max(M.sum(), 1e-9)
            v[tname] = dict(spearman=spearman(m / A[ok], tgt / A[ok]),
                            moved_share=float(0.5 * np.abs(m / m.sum() - tgt / tgt.sum()).sum()))
        res["stages"][k] = v
        print(f"  {k:54s} towns: rank {v['towns']['spearman']:.3f} move {v['towns']['moved_share']:.1%}   "
              f"area: rank {v['town area']['spearman']:.3f} move {v['town area']['moved_share']:.1%}")
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(res, indent=1))


if __name__ == "__main__":
    main()
